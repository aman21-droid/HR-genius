"""Build the editable HRGenius presentation from local project evidence."""
from pathlib import Path
import os, sys, re, json
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'hrgenius-ppt-tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'reports'
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
C = dict(bg='F7F7F2', white='FFFFFF', ink='2D3424', olive='5C6936', deep='293121', muted='626752', pale='ECEFDC', line='E3E5D8', lime='CCD79E', gold='DAB66A', blue='4D657E')
slides = [prs.slides.add_slide(prs.slide_layouts[6]) for _ in range(14)]
notes = []

def rgb(c): return RGBColor.from_string(C.get(c, c))
def rect(s,x,y,w,h,fill='white',radius=True,line=None):
    z=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    z.fill.solid(); z.fill.fore_color.rgb=rgb(fill)
    if line: z.line.color.rgb=rgb(line)
    else: z.line.fill.background()
    if radius: z.adjustments[0]=.13
    return z
def txt(s,x,y,w,h,t,size=18,color='ink',bold=False,font='Segoe UI'):
    z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    f=z.text_frame; f.word_wrap=True
    f.margin_left=f.margin_right=0; f.margin_top=f.margin_bottom=0
    for i,l in enumerate(t.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=l; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.space_after=Pt(8)
    return z
def button(s,x,y,w,label,target=None,url=None,fill='olive',color='white'):
    z=rect(s,x,y,w,.44,fill)
    z.text_frame.text=label; z.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    z.text_frame.margin_top=z.text_frame.margin_bottom=0
    p=z.text_frame.paragraphs[0]; p.alignment=1; p.font.name='Segoe UI'; p.font.size=Pt(12); p.font.bold=True; p.font.color.rgb=rgb(color)
    if target is not None: z.click_action.target_slide=slides[target-1]
    if url: z.click_action.hyperlink.address=url
    return z
def pic(s,name,x,y,w,h=None):
    path=ROOT/name
    from PIL import Image
    iw,ih=Image.open(path).size
    if h is None: h=w*ih/iw
    # Contain, never stretch the screenshot.
    scale=min(w/iw,h/ih); aw,ah=iw*scale,ih*scale
    return s.shapes.add_picture(str(path),Inches(x+(w-aw)/2),Inches(y+(h-ah)/2),width=Inches(aw),height=Inches(ah))
def base(n,kicker,title,subtitle='',dark=False):
    s=slides[n-1]; s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('deep' if dark else 'bg')
    txt(s,.55,.3,9,.25,'HRGenius  /  '+kicker.upper(),10,'lime' if dark else 'olive',True)
    txt(s,.55,.88,12.1,.7,title,32,'white' if dark else 'ink',True)
    if subtitle: txt(s,.58,1.65,12,.58,subtitle,15,'lime' if dark else 'muted')
    rect(s,.55,6.96,12.2,.012,'olive' if dark else 'line',False)
    txt(s,.55,7.13,7,.2,'A LITTLE MORE HUMAN.   •   '+('OPTIONAL DEEP DIVE' if n>10 else 'PROJECT SHOWCASE'),9,'lime' if dark else 'muted')
    button(s,10.15,7.04,.85,'Menu',2,fill='pale',color='ink')
    button(s,11.1,7.04,.42,'‹',max(1,n-1),fill='pale',color='ink')
    button(s,11.62,7.04,.42,'›',min(14,n+1),fill='pale',color='ink')
    txt(s,12.25,7.12,.6,.2,f'{n:02}',10,'lime' if dark else 'muted',True)
    trans=OxmlElement('p:transition'); trans.set('spd','med'); trans.append(OxmlElement('p:fade')); s._element.append(trans)
    return s
def note(n,title,body):
    slides[n-1].notes_slide.notes_text_frame.text=title+'\n\n'+body
    notes.append((n,title,body))
def card(s,x,y,w,h,number,title,body,fill='white',target=None):
    z=rect(s,x,y,w,h,fill)
    if target: z.click_action.target_slide=slides[target-1]
    txt(s,x+.22,y+.18,w-.44,.28,number,11,'olive',True)
    t=txt(s,x+.22,y+.64,w-.44,.54,title,21,'ink',True)
    b=txt(s,x+.22,y+1.25,w-.44,h-1.35,body,15,'muted')
    if target:
        t.click_action.target_slide=slides[target-1]; b.click_action.target_slide=slides[target-1]
def line(s,x1,y1,x2,y2,color='olive'):
    z=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); z.line.color.rgb=rgb(color); z.line.width=Pt(2)

# 01 — Cover
s=base(1,'Human resource management','',dark=True)
pic(s,'frontend/src/assets/olive-workspace.png',6.9,.75,5.9,5.9)
txt(s,.7,1.2,6,.6,'HRGenius',40,'white',True)
txt(s,.7,2.1,6,1.8,'Good work starts\nwith happy people.',35,'white',False,'Georgia')
txt(s,.73,4.18,5.4,.8,'One connected workplace for people,\noperations and everyday appreciation.',19,'lime')
button(s,.73,5.48,2.3,'Explore the story  →',2,fill='lime',color='deep')
txt(s,3.3,5.59,3.2,.5,'10 core slides · 4 optional dives',11,'lime')
note(1,'Opening / 20 seconds',"HRGenius is a human resource management system that connects employee information, daily self-service, operational workflows and workforce insight. Open with a relatable question: how many places would an employee normally visit to request leave, find a payslip and thank a teammate? This project puts those experiences into one permission-aware workspace. The presentation follows a roughly six-to-eight-minute core story; technical detail is available through four optional slides and these notes. All product screenshots come from the local HRGenius demonstration environment. Their names and figures are seeded demonstration data, not customer results. Use Slide Show mode (F5) and click Explore the story. The deck is editable and does not use macros.")

# 02 — Interactive hub
s=base(2,'Choose your route','A presentation you can explore.','Click a card for the part you care about, or follow the arrows for the complete story.')
for x,y,num,title,body,target in [(.6,2.4,'01 / WHY','The big idea','The problem, the people,\nand the purpose.',3),(4.72,2.4,'02 / EXPERIENCE','Feel the product','Olive and white, thoughtful\ndetails, real interaction.',4),(8.84,2.4,'03 / WORKFLOWS','Follow the journey','From a first application\nto everyday growth.',5),(.6,4.58,'04 / ENGINEERING','Under the surface','Architecture, permissions\nand sensitive information.',7),(4.72,4.58,'05 / EVIDENCE','See the checks','What was verified, and\nwhat still needs validation.',9),(8.84,4.58,'06 / YOUR TURN','Pick a mini-demo','Explore a chart, recognize\na colleague, switch themes.',10)]:
    card(s,x,y,3.88,2.08,num,title,body,target=target)
note(2,'How to navigate',"This is the main navigation hub. The six cards are actual internal PowerPoint hyperlinks. Menu returns here from every slide. The left and right arrows provide a conventional sequence. For a short product pitch use slides 1–6 and 10; for a technical project review continue through 7–9 and use appendices 12–14 for questions. Slide 4 links to a dark-mode comparison on slide 11. Slide 10 offers offline slide tours plus explicit external links to the running application. External links require the local application and a suitable signed-in account; the screenshots remain available offline. There is no embedded working HR application inside the PowerPoint.")

# 03 — Problem / solution
s=base(3,'The purpose','Less chasing. More clarity.','A shared workspace for employees, managers and HR operations.')
rect(s,.6,2.45,4.0,3.8,'deep')
txt(s,.9,2.77,3.4,.3,'THE EVERYDAY FRICTION',11,'lime',True)
txt(s,.9,3.42,3.35,1.8,'Scattered records.\nUnclear ownership.\nRepeated follow-ups.',25,'white')
txt(s,.9,5.53,3.3,.4,'A familiar problem worth solving.',14,'lime')
for y,a,b in [(2.5,'One employee context','Profiles, documents and organization structure.'),(3.75,'A visible next step','Requests, approvals and personal tasks.'),(5.,'A more human routine','Self-service, recognition and readable insights.')]:
    rect(s,4.92,y,7.8,1.08,'white')
    txt(s,5.2,y+.15,7.1,.35,a,20,'olive',True); txt(s,5.2,y+.6,7.1,.35,b,15,'muted')
note(3,'Problem, objectives and users',"The project addresses fragmentation in employee records and HR processes. Objectives are to provide a consistent employee source of information, support traceable request workflows, enable self-service, preserve role-appropriate access, and make operational data easier to understand. The intended users include employees, managers, recruiters, payroll administrators, HR managers, HR administrators and super administrators. Employees access their own information and applicable self-service features; operational specialists receive permissions for their responsibilities. Managers have team-oriented responsibilities where the backend permits them. Benefits on this slide are design goals and qualitative outcomes, not measured cost savings or adoption results. Do not claim percentage improvements in productivity or processing time without a separate study. The system includes public careers pages outside the authenticated application shell.")

# 04 — UI light
s=base(4,'The experience','Calm on the surface. Useful in the details.','A real dashboard from the project’s demonstration environment.')
rect(s,.6,2.3,8.1,4.35,'white')
pic(s,'reports/ui-dashboard-light.png',.68,2.35,7.95,4.23)
txt(s,9.02,2.5,3.55,.65,'White + olive',25,'olive',True)
txt(s,9.02,3.24,3.55,2.2,'A warm daily greeting\nReadable cards and actions\nCharts with a data view\nKeyboard page search\nResponsive navigation',17)
for i,c in enumerate(['white','pale','olive','deep']): rect(s,9.02+i*.6,5.51,.43,.32,c)
button(s,9.02,6.06,3.48,'Switch to the dark-side view  →',11)
note(4,'Design system and dashboard tour',"The light theme uses white surfaces, a warm off-white background (#F7F7F2), olive actions (#5C6936), a deep olive sidebar (#293121) and dark readable text (#2D3424). The screenshot is a real captured dashboard; the counts and financial figure are demonstration records. The greeting changes with time of day. Permissions influence visible actions and access to analytics. The implemented workforce chart supports six- and twelve-month periods, keyboard/hover values and a data-table alternative. Department bars and policy acknowledgement progress summarize existing API data. The shared shell includes Ctrl/Cmd+K page search, arrow-key navigation, Enter to select, Escape to close, a mobile drawer, a skip link and a persistent theme preference. First visit defaults to the light theme. Dark mode remains an intentional user option. Do not claim that individual chart points already drill into filtered business lists: that enhancement was discussed but is not part of the verified implementation. Click the theme comparison button for the corresponding dark screenshot.")

# 05 — Full module coverage
s=base(5,'The employee journey','From “you’re hired” to “thank you”.','Connected modules cover the employee lifecycle and the work around it.')
items=[('01','Attract & welcome','Careers · candidates · interviews\nOffers · onboarding plans'),('02','Know your people','Profiles · org chart · documents\nOrganization masters · assets'),('03','Manage the day','Attendance · regularization\nLeave balances · approvals'),('04','Pay with structure','Salary components · payroll runs\nAdjustments · payslips · exports'),('05','Grow & recognize','Review cycles · goals · progress\nSelf / manager reviews · praise'),('06','Support & understand','Helpdesk · policies · acknowledgements\nAnalytics · personal summary')]
for i,(a,b,c) in enumerate(items): card(s,.6+(i%3)*4.13,2.4+(i//3)*2.16,3.89,2.03,a,b,c)
button(s,9.5,6.62,3.15,'Open the technical catalogue  →',12,fill='pale',color='ink')
note(5,'Complete functional scope',"Recruitment includes requisitions, public careers, candidates, applications and pipeline events, interview scheduling/feedback, offers and onboarding. Onboarding provides templates, plans and tasks. Core HR includes employee records, contacts, statutory details, document metadata and file access, reporting relationships, timeline events, organization masters and asset allocation history. Time management includes attendance days, holidays, regularization, leave types, balances, requests, cancellations and accrual operations, integrated with approval requests and steps. Payroll includes salary components, payroll runs, adjustments, calculation, submission, paid state, payslips, PDF downloads and register/bank-file exports. Performance includes review cycles, goals, progress updates, employee self-review, manager review, acknowledgement and feedback/praise. Support includes helpdesk tickets/comments and policies/acknowledgements. Analytics and personal summaries assemble existing operational information. This scope was checked against current frontend routes, backend controllers and Flyway migrations; the older project report describes an earlier milestone and is not the sole source. Module presence does not itself prove every business scenario has been tested.")

# 06 — human interaction / workflows
s=base(6,'Small moments, clear actions','Make routine work feel considered.','A request has a next step. A contribution gets a moment of appreciation.')
for i,(a,b) in enumerate([('Request','Employee applies'),('Review','Assigned approver acts'),('Outcome','Status becomes visible')]):
    x=.65+i*2.62
    rect(s,x,2.55,2.38,1.25,'pale')
    txt(s,x+.18,2.74,2.05,.36,a,22,'olive',True); txt(s,x+.18,3.25,2.05,.35,b,13)
    if i<2: txt(s,x+2.4,2.98,.22,.3,'›',20,'olive',True)
txt(s,.7,4.2,6.7,.5,'Interaction with a purpose',24,'ink',True)
txt(s,.7,4.94,6.5,1.25,'Select a day to inspect attendance.\nCompare chart periods and read exact values.\nCelebrate someone through the praise wall.',19)
rect(s,8.66,2.4,4.03,4.22,'white')
pic(s,'reports/ui-appreciation-dialog.png',8.83,2.55,3.68,3.3)
txt(s,8.95,6.0,3.5,.4,'A real appreciation form',14,'olive',True)
note(6,'Workflow and interaction details',"Use a leave request as the simple explanatory workflow: an employee submits a request, the authorized approval process handles it, and the visible outcome updates. This is a conceptual summary; the configured workflow and policy determine the actual approval stages and eligibility. Attendance adds calendar-day selection with check-in, check-out and worked-time detail; changing month clears the selected day. Leave balance graphics separate used, pending and available amounts. Payroll comparison uses existing recent-run data. The appreciation wall reads existing feedback and offers a Give shoutout action using the existing feedback form and service. The carousel and dialog were tested without posting a recognition message or making a business-data mutation. Other details include loading, empty and error states, retries, accessible chart values, responsive layouts, non-nested org-chart controls and reduced-motion behavior. These are frontend presentation improvements built around existing APIs. Avoid inventing birthdays, employee anniversaries or recognition messages when no corresponding data is available.")

# 07 — Architecture
s=base(7,'Engineering','One interface. Clear layers underneath.','A modular application with a REST boundary and server-enforced access.')
for x,k,title,body in [(.65,'PRESENTATION','Angular 18','Standalone components\nMaterial + SCSS\nLazy routes · guards · signals'),(4.78,'APPLICATION','Spring Boot 3.3','Java 21 target\nController → service → repository\nSecurity · validation · JPA'),(8.9,'PERSISTENCE','Oracle / H2 demo','Oracle target deployment\nH2 development profile\nFlyway schema migrations')]:
    card(s,x,2.6,3.77,2.7,k,title,body,fill='white')
txt(s,4.42,3.65,.35,.4,'→',24,'olive',True); txt(s,8.54,3.65,.35,.4,'→',24,'olive',True)
rect(s,.65,5.63,12.02,.75,'pale')
txt(s,.94,5.85,11.5,.35,'REST /api/v1   •   DTOs + MapStruct   •   File storage   •   Mail notifications   •   OpenAPI',16,'olive',True)
note(7,'Architecture and technology rationale',"The frontend uses Angular 18 standalone components with lazy-loaded feature routes, Angular Material and SCSS. Route guards provide navigation behavior; the backend is responsible for authorization. The backend targets Java 21 and Spring Boot 3.3.4. Features are organized into packages with controllers, services, repositories and DTOs; JPA/Hibernate persists entities, MapStruct maps data and Bean Validation validates inputs. Spring Security and JWT protect REST endpoints under /api/v1. Oracle is the intended database for the packaged deployment. The running demonstration used an H2 profile with Oracle compatibility; it must not be described as proof of production Oracle behavior. Flyway manages versioned schema changes. Document metadata lives in the database while file bytes are handled by a storage service. The codebase also contains notification/mail handling, Caffeine caching and springdoc OpenAPI support. The architecture is a modular backend application, not a deployed microservice fleet. Keeping the existing REST and business-service boundary let the UI redesign change presentation without rewriting backend contracts.")

# 08 — Security
s=base(8,'Trust and access','The right information, for the right person.','Permissions and record-level rules protect more than the navigation menu.')
card(s,.65,2.48,3.78,3.72,'01 / IDENTITY','Authenticate','JWT access + refresh flow\nBCrypt password hashing\nRefresh-token rotation\nFailed-login lockout',fill='white')
card(s,4.78,2.48,3.78,3.72,'02 / AUTHORIZATION','Check the scope','Role and permission checks\nSelf / team / HR boundaries\nBackend authorization\nAudit history',fill='white')
card(s,8.9,2.48,3.78,3.72,'03 / SENSITIVE DATA','Reveal with care','Selected fields encrypted\nMasked statutory / bank data\nControlled reveal actions\nValidated file handling',fill='white')
txt(s,.7,6.45,12,.3,'Employee  ·  Manager  ·  Recruiter  ·  Payroll Admin  ·  HR Manager  ·  HR Admin  ·  Super Admin',13,'olive',True)
note(8,'Security boundaries and important limits',"The access model combines roles, permissions and record-level scope. An employee's self-service access differs from manager team access and authorized HR operations; specific permissions are checked by endpoint and service. JWT access tokens and rotating refresh tokens are used, with BCrypt password hashes and a failed-login lockout policy. Logout revokes refresh-token access; do not claim every issued access token is instantly invalidated. Selected statutory and bank fields use AES-GCM at rest with masked responses and controlled reveal paths; this does not mean every database column is encrypted. Sensitive actions and operational changes have audit support. Upload handling uses generated storage keys, allowed file types/signature validation and validated paths. It is not a malware-scanning service. HTTPS, managed secrets, deployment hardening, recovery procedures and security testing remain deployment responsibilities. The slide describes implemented controls, not a compliance certification or penetration-test result. Permission-filtered UI improves usability, but a hidden button is never the security boundary.")

# 09 — Evidence
s=base(9,'Validation','Beautiful is better when it behaves.','Recorded frontend checks from this workspace—not a production certification.')
for x,num,label,detail in [(.65,'36','page / theme scans','18 routes across light and dark'),(4.78,'19','interaction checks','Keyboard, charts, dialogs and more'),(8.9,'18','mobile routes','390 px viewport layout checks')]:
    rect(s,x,2.55,3.78,2.53,'white'); txt(s,x+.25,2.77,3.1,.95,num,52,'olive',True)
    txt(s,x+.25,3.91,3.2,.4,label,20,'ink',True); txt(s,x+.25,4.49,3.23,.4,detail,12,'muted')
rect(s,.65,5.45,7.72,.97,'pale'); txt(s,.9,5.65,7.2,.65,'Production frontend build passed.\nAutomated accessibility scans reported zero violations.',16,'olive',True)
txt(s,8.95,5.52,3.65,.9,'Still to validate: production Oracle, end-to-end business flows, load and recovery.',15,'muted')
note(9,'Verification evidence and limits',"Evidence is stored in reports/ui-audit.json, ui-interactions.json, ui-mobile.json and ui-final.json. The audit records 18 routes in each of two themes, totaling 36 automated accessibility scans, with zero reported violations and no recorded page runtime errors. Nineteen interaction checks passed, including dashboard period selection, table rows, keyboard chart values, appreciation carousel and dialog cancellation, keyboard search, saved theme, attendance selection/reset, payroll tooltip, analytics period/table, org-chart keyboard behavior, mobile drawer behavior, analytics retry and employee permission restrictions. Eighteen mobile routes were checked at a 390-pixel viewport; native tab strips can scroll internally. Final checks covered first-visit light theme, light/dark dashboard scans and graceful fallback for failed organization lookups. The production Angular build passed within configured budgets. Backend test classes exist for auth, employee, leave/attendance, recruitment, payroll, performance, services, migrations and personal summary, but were not rerun for this presentation. Automated accessibility scans do not replace manual assistive-technology testing. No load benchmark, production Oracle acceptance, backup-restore exercise or business-result improvement is claimed.")

# 10 — Demo & close
s=base(10,'Your turn','Pick a 30-second moment.','Choose an offline tour, or open the running project for a live demonstration.',dark=True)
for x,num,title,body,target,path in [(.65,'01 / EXPLORE','Find the signal','Change a chart period.\nRead the exact values.',4,'analytics'),(4.78,'02 / APPRECIATE','Make someone’s day','Open the praise wall.\nShow the shoutout form.',6,'performance'),(8.9,'03 / PERSONALIZE','Set the mood','Compare olive + white\nwith the dark theme.',11,'dashboard')]:
    card(s,x,2.5,3.78,2.95,num,title,body,fill='pale',target=target)
    button(s,x,5.67,1.81,'Slide tour  →',target,fill='lime',color='deep')
    button(s,x+1.98,5.67,1.8,'Open app  ↗',url='http://localhost:4200/'+path,fill='white',color='deep')
txt(s,.7,6.38,12,.35,'Live links need the local app and an authorized account. Offline slide tours work without a server.',12,'lime')
note(10,'Demo script, conclusion and questions',"Invite the audience to choose one moment. Explore: open Analytics with an authorized account, change between six and twelve months, inspect a data point by keyboard or pointer and show the tabular alternative. Appreciate: open Performance or the dashboard praise wall, show the carousel and open Give shoutout, then cancel unless posting is explicitly intended. Personalize: use the dashboard theme control and show its persistence and readable actions. The presentation's Slide tour buttons are offline internal links. Open app buttons use localhost:4200 and require the frontend, its configured backend and an appropriate session; on another computer localhost means that computer. Access may be denied or a login page shown if the account lacks permission. Close the core presentation here: HRGenius combines employee operations with a calm, accessible and human-centered interface. Ask which daily HR interaction the audience would improve next. Optional material follows: dark-mode comparison, module/API catalogue, logical data model and deployment/future work. No real business data needs to be created during this demo.")

# 11 — dark comparison
s=base(11,'Theme comparison','Same workspace. A softer evening view.','Dark mode preserves hierarchy, contrast and a visible way back to light.',dark=True)
pic(s,'reports/ui-dashboard-dark.png',.65,2.35,8.15,4.3)
txt(s,9.02,2.63,3.6,.5,'Intentional contrast',23,'white',True)
txt(s,9.02,3.44,3.45,1.8,'Deep olive surfaces\nPale olive primary actions\nReadable labels and borders\nA remembered preference',18,'lime')
button(s,9.02,5.76,3.5,'Return to olive + white  →',4,fill='lime',color='deep')
note(11,'Optional / dark theme',"This screenshot shows the implemented dark theme, not a different product. Its background is #191D14, surfaces #252B1E, primary action color #CCD79E and primary text #F0F2E8. The selected sidebar state remains light and visible. Controls have separate surface, border, foreground, hover and focus treatments rather than merely darkening an entire screenshot. The preference persists between reloads, while a new session with no saved preference starts light even when the operating system prefers dark. The app has a visible Light mode / Dark mode control. Automated checks include the principal routes in both themes; automated success does not prove every contrast state or every assistive-technology interaction has been manually reviewed. Click Return to olive + white to compare with slide 4, or Menu to choose another chapter.")

# 12 — API detail
s=base(12,'Functional + API reference','The complete scope, grouped for review.','All API roots below are prefixed with /api/v1. Full method contracts live in the controllers / OpenAPI.')
groups=[('PEOPLE & ORGANIZATION','/employees · /org · /org-chart\n/documents · /assets · /me','Profiles, statutory data, files, reporting lines,\nmasters, asset assignments and personal summary.'),('TIME & APPROVALS','/attendance · /holidays\n/leave · /approvals','Daily records, regularization, calendars, balances,\nrequests, accrual and approval steps.'),('TALENT & PAY','/recruitment · /public/careers\n/onboarding · /payroll','Applications, interviews, offers, starter tasks,\ncomponents, runs, adjustments and payslips.'),('GROWTH, SUPPORT & CONTROL','/performance · /helpdesk · /policies\n/analytics · /audit-logs · /auth','Reviews, goals, praise, tickets, acknowledgements,\ninsight, audit history and authentication.')]
for i,(a,b,c) in enumerate(groups):
    x=.65+(i%2)*6.17;y=2.35+(i//2)*2.18
    rect(s,x,y,5.94,1.99,'white');txt(s,x+.22,y+.16,5.5,.3,a,11,'olive',True);txt(s,x+.22,y+.56,5.5,.6,b,16,'ink',True);txt(s,x+.22,y+1.3,5.5,.58,c,14,'muted')
note(12,'Detailed feature and API catalogue',"Core HR roots: /employees for directory and employee lifecycle, nested /employees/{employeeId} for profile subresources, /org for organization masters, /org-chart for reporting structure, /documents for document operations, /assets for allocation management and /me for personal summary. Time roots: /attendance, /holidays, /leave and /approvals. Leave exposes types, balances, adjustments, requests, cancellation and accrual operations. Recruitment roots: /recruitment and /public/careers; onboarding root: /onboarding. Payroll: /payroll/runs supports listing/detail, create, calculate, submit, paid state and deletion as permitted; nested run endpoints expose payslips, register, bank file and adjustments. /payroll/components manages salary components; /payroll/me/payslips, /me/structure and /payslips/{id}/pdf support self-service and downloads. Performance: /performance/cycles includes launch, close, summary and review lists; reviews support goals, progress, self/manager assessment and acknowledgement; feedback has wall, received, given and team views. Support/control roots are /helpdesk, /policies, /analytics, /audit-logs and /auth. These are API groupings, not a claim that every action is public or available to every role. Check controller annotations and services for exact HTTP methods, parameters, validation and permissions. Source: backend/src/main/java/com/hrgenius/**/controller and frontend/src/app/app.routes.ts.")

# 13 — data model
s=base(13,'Data model','An employee-centered model with history.','Selected logical relationships; the migrations contain the full physical schema.')
for x,y,w,h,head,body in [(4.6,3.47,4.1,1.08,'EMPLOYEES','Identity · role in organization · manager'),(.65,2.48,3.5,1.15,'IDENTITY & ACCESS','Users ↔ roles ↔ permissions\nRefresh tokens · audit log'),(9.18,2.48,3.5,1.15,'ORGANIZATION','Department · designation · grade\nLocation · unit · cost center'),(.65,4.91,3.5,1.3,'DAILY WORK','Attendance · leave · approvals\nDocuments · asset assignments'),(9.18,4.91,3.5,1.3,'PAY & GROWTH','Payslips → lines; payroll runs\nReviews → goals; feedback')]:
    rect(s,x,y,w,h,'pale' if head=='EMPLOYEES' else 'white');txt(s,x+.18,y+.14,w-.36,.3,head,12,'olive',True);txt(s,x+.18,y+.54,w-.36,.65,body,13)
for p in [(4.15,3.05,4.6,3.85),(9.18,3.05,8.7,3.85),(4.15,5.47,4.6,4.13),(9.18,5.47,8.7,4.13)]: line(s,*p)
txt(s,4.61,2.48,4.1,.65,'Candidate → application → offer\nOnboarding plan → tasks',15,'muted')
line(s,6.65,3.16,6.65,3.47)
txt(s,4.65,5.2,4.0,.85,'Also linked: helpdesk tickets,\npolicy acknowledgements\nand employee timeline events.',14,'muted')
note(13,'Schema, relationships and persistence',"The diagram is intentionally a selected logical model, not an exhaustive ERD or a statement of every foreign-key cardinality. Flyway migrations define authentication tables (users, roles, permissions and join tables), refresh tokens and audit_log; organization masters; employees and their related statutory, emergency contact, document and timeline records; assets and assignment history; approval requests/steps; leave types, balances, requests and accrual log; holidays, attendance days and regularization requests. Recruitment tables include job_requisitions, candidates, job_applications, application_events, interviews, interview_feedback and offers. Onboarding includes templates/template tasks and plans/tasks. Payroll includes salary_components, payroll_runs, payslips, payslip_lines and payroll_adjustments. Performance includes review_cycles, performance_reviews, goals and feedback_notes. Support/compliance includes helpdesk_tickets, helpdesk_comments, policies and policy_acknowledgements. Employees have organization references and a manager self-reference. Users may link to an employee; do not imply that employee_id is necessarily a unique schema constraint without checking the migration. File metadata and storage bytes are separate. Versioned migrations document evolution of the database, but the H2 demo is not a substitute for Oracle integration and migration testing.")

# 14 — delivery and future
s=base(14,'Delivery + next steps','Ready to demonstrate. Room to grow.','Keep implementation, verification and future work distinct.')
card(s,.65,2.5,3.78,3.65,'RUN THE PROJECT','Development & deploy','Angular frontend\nSpring Boot REST backend\nOracle deployment / H2 demo\nDocker Compose configuration\nFile storage + mail service')
card(s,4.78,2.5,3.78,3.65,'CHECK THE PROJECT','Quality evidence','Frontend build + UI checks\nBackend test sources present\nOracle acceptance still needed\nBusiness-flow regression\nManual accessibility review')
card(s,8.9,2.5,3.78,3.65,'EVOLVE THE PROJECT','Proposed next steps','Deeper role-based dashboards\nChart-to-record navigation\nSSO and stronger operations\nMonitoring + recovery drills\nMeasured user feedback')
button(s,.65,6.4,3.0,'Back to the demo choices  →',10)
note(14,'Deployment, limitations, future scope and sources',"The repository provides an Angular frontend, a Spring Boot backend and a Docker Compose deployment configuration with Oracle, frontend serving, backend and mail development tooling. The current screenshots came from localhost:4200 with the development backend; environments may use different configured ports. H2 in-memory demonstration data can reset when the backend restarts. Production needs persistent database and upload storage, externalized secrets, HTTPS, backups and restore verification, monitoring and deployment-specific configuration. Existing backend tests cover several modules, but no backend test execution is represented as newly completed in this presentation. Next validation priorities are Oracle integration/migration acceptance, complete business-flow regression, manual accessibility testing, security review, realistic load tests and recovery exercises. Proposed product enhancements include richer role-specific dashboard composition and chart-to-record navigation; they were discussed but should not be presented as completed. SSO and operational hardening are future directions rather than current verified features. Primary sources: frontend/src/app/app.routes.ts; frontend/src/app/features and shared components; backend/pom.xml; backend/src/main/java/com/hrgenius; backend/src/main/resources/db/migration; reports/UI_DESIGN_NOTES.md; reports/ui-audit.json; reports/ui-interactions.json; reports/ui-mobile.json; reports/ui-final.json. The older Project Report and SRS are supporting documents; current source code takes precedence where scope differs. Slides and notes contain no real credentials, macros or unsupported business benefit statistics.")

prs.core_properties.title='HRGenius — A Little More Human'
prs.core_properties.subject='Interactive project showcase with technical presenter notes'
prs.core_properties.author='HRGenius Project'
prs.core_properties.keywords='HRGenius, HRMS, olive, white, interactive, project'
prs.save(OUT/'HRGenius_Interactive_Showcase.pptx')
guide=['# HRGenius — presentation and presenter notes','',
'Open **HRGenius_Interactive_Showcase.pptx** in PowerPoint and press F5. Cards and navigation buttons are clickable in Slide Show mode. Use Presenter View for notes. Core story: slides 1–10 (about 6–8 minutes); optional detail: slides 11–14. The PDF is a static reading copy, not the interactive presentation.','',
'The internal slide tours work offline. “Open app” links require HRGenius running on localhost:4200 and a suitably authorized login. They open the browser and do not submit data. On a different computer, localhost refers to that computer.','']
for n,t,b in notes: guide += [f'## {n:02}. {t}','',b,'']
(OUT/'HRGenius_Presenter_Notes.md').write_text('\n'.join(guide),encoding='utf-8')
print(json.dumps({'slides':len(slides),'pptx':str(OUT/'HRGenius_Interactive_Showcase.pptx'),'size_bytes':(OUT/'HRGenius_Interactive_Showcase.pptx').stat().st_size}))
